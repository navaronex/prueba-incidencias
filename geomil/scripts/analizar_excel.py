"""Inspección de solo lectura. No importa ni corrige registros automáticamente."""
import argparse
from collections import Counter
from datetime import datetime, date
import hashlib
import json
from pathlib import Path
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('archivo', type=Path)
parser.add_argument('--salida', type=Path, required=True)
args = parser.parse_args()
wb = openpyxl.load_workbook(args.archivo, data_only=False)
sheet = wb['PERDIDA MENSAJERIAS']
def clean(value):
    return '' if value is None else str(value).strip()
rows = [row for row in sheet.iter_rows(min_row=8) if clean(row[1].value)]
operations = Counter(clean(row[1].value) for row in rows)
report = {
    'archivo': args.archivo.name,
    'sha256': hashlib.sha256(args.archivo.read_bytes()).hexdigest(),
    'criterio_registro': 'Fila desde la 8 con número de operación no vacío en B. Son candidatos, no expedientes validados.',
    'hojas': [],
    'candidatos_principal': len(rows),
    'operaciones_distintas': len(operations),
    'operaciones_repetidas': {k: v for k,v in operations.items() if v > 1},
    'categorias_originales': {
        label: dict(Counter(clean(row[col].value) or '(vacío)' for row in rows))
        for label,col in [('mensajeria',5),('tipo',9),('estado',10)]
    },
    'campos_vacios': {
        label: [row[col].coordinate for row in rows if not clean(row[col].value)]
        for label,col in [('expedicion',0),('cliente',3),('fecha_envio',7),('fecha_incidencia',8),('estado',10),('responsable',11)]
    },
    'fechas_texto': [dict(celda=c.coordinate, valor=c.value) for row in rows for c in row[7:9] if isinstance(c.value,str)],
    'fechas_invertidas': [row[0].row for row in rows if isinstance(row[7].value,(datetime,date)) and isinstance(row[8].value,(datetime,date)) and row[8].value < row[7].value],
}
for s in wb:
    formulas = [(c.coordinate,c.value) for row in s for c in row if c.data_type == 'f']
    report['hojas'].append({
        'nombre':s.title, 'filas_utilizadas':s.max_row, 'columnas_utilizadas':s.max_column,
        'formulas':formulas, 'combinadas':[str(r) for r in s.merged_cells.ranges],
        'referencias_entre_hojas':[coord for coord,f in formulas if '!' in f],
    })
args.salida.parent.mkdir(parents=True,exist_ok=True)
args.salida.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
